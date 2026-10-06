import calendar
from datetime import datetime
import io
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import streamlit as st
from supabase import create_client

# Danh sách ký hiệu chấm công hợp lệ chuẩn y tế / cơ quan
VALID_CODES_MAP = {
    "X": "X",
    "1": "X",
    "+": "X",
    "P": "P",
    "OM": "ÔM",
    "ÔM": "ÔM",
    "CO": "CO",
    "CT": "CT",
    "T": "T",
    "TS": "TS",
    "KL": "KL",
    "R": "R",
    "H": "H",
    "L": "L",
    "F": "F",
    "": "",
}


# ---------------------------------------------------------
# 1. KẾT NỐI SUPABASE
# ---------------------------------------------------------
@st.cache_resource
def init_supabase():
    try:
        url = st.secrets.get("SUPABASE_URL", "").strip()
        key = (
            st.secrets.get("SUPABASE_KEY")
            or st.secrets.get("SUPABASE_SERVICE_KEY", "")
        ).strip()
        if not url or not key:
            return None
        if not url.startswith("https://"):
            url = f"https://{url}"
        return create_client(url, key)
    except Exception:
        return None


# ---------------------------------------------------------
# 2. KIỂM TRA NGÀY LỄ VIỆT NAM & CHUẨN HÓA CHUỖI
# ---------------------------------------------------------
def is_vietnam_holiday(day, month, year=2026):
    """Xác định các ngày nghỉ lễ chính thức theo quy định"""
    fixed_holidays = [
        (1, 1),  # Tết Dương lịch
        (30, 4),  # 30/4 Giải phóng miền Nam
        (1, 5),  # 1/5 Quốc tế Lao động
        (1, 9),  # 1/9 Nghỉ Quốc khánh
        (2, 9),  # 2/9 Quốc khánh
    ]
    if (day, month) in fixed_holidays:
        return True

    if year == 2026:
        lunar_holidays = [
            (16, 2),
            (17, 2),
            (18, 2),
            (19, 2),
            (20, 2),
            (21, 2),
            (26, 4),
        ]
        if (day, month) in lunar_holidays:
            return True

    return False


def normalize_str(s):
    """Chuẩn hóa chuỗi ký tự để so sánh đơn vị chính xác 100%"""
    if pd.isna(s) or s is None:
        return ""
    s_str = str(s).strip()
    if s_str.lower() in ["nan", "none", "null", "<na>", ""]:
        return ""
    return " ".join(s_str.lower().split())


def clean_and_uppercase_code(val):
    """Tự động chuyển chữ thường sang in hoa & chuẩn hóa ký hiệu chấm công"""
    if pd.isna(val) or val is None:
        return ""
    v_str = str(val).strip().upper()
    if v_str in ["NAN", "NONE", "<NA>", "0", "NULL"]:
        return ""
    return VALID_CODES_MAP.get(v_str, v_str)


# ---------------------------------------------------------
# 3. QUÉT TỰ ĐỘNG DANH SÁCH ĐƠN VỊ TỪ CSDL
# ---------------------------------------------------------
def load_system_units(supabase_client):
    """Gom toàn bộ Đơn vị khai báo từ 'Danh mục Hệ thống' VÀ 'Hồ sơ Cán bộ'"""
    units_set = set()

    if supabase_client:
        try:
            res = (
                supabase_client.table("categories")
                .select("content")
                .eq("cat_key", "don_vi")
                .execute()
            )
            if res.data and res.data[0].get("content"):
                df_dv = pd.DataFrame(res.data[0]["content"])
                for col in df_dv.columns:
                    c_lower = str(col).strip().lower()
                    if c_lower not in [
                        "stt",
                        "mã",
                        "mã đv",
                        "mã đơn vị",
                        "id",
                        "code",
                    ]:
                        vals = (
                            df_dv[col].dropna().astype(str).str.strip().tolist()
                        )
                        for v in vals:
                            if v and v.lower() not in ["nan", "none", ""]:
                                units_set.add(v)
        except Exception:
            pass

    df_emp = pd.DataFrame()
    if "employees_profile" in st.session_state:
        df_emp = st.session_state["employees_profile"]
    elif supabase_client:
        try:
            res_emp = (
                supabase_client.table("categories")
                .select("content")
                .eq("cat_key", "employees_profile")
                .execute()
            )
            if res_emp.data and res_emp.data[0].get("content"):
                df_emp = pd.DataFrame(res_emp.data[0]["content"])
        except Exception:
            pass

    if isinstance(df_emp, pd.DataFrame) and not df_emp.empty:
        for col in df_emp.columns:
            c_lower = str(col).strip().lower()
            if any(
                k in c_lower
                for k in ["đơn vị", "khoa", "phòng", "bộ phận", "trung tâm"]
            ):
                vals = df_emp[col].dropna().astype(str).str.strip().tolist()
                for v in vals:
                    if v and v.lower() not in ["nan", "none", ""]:
                        units_set.add(v)

    if not units_set:
        return [
            "Ban Giám đốc",
            "Khoa Cấp cứu",
            "Khoa Ngoại Tổng hợp",
            "Phòng Điều dưỡng",
            "Phòng Nhân sự - Tổng hợp",
            "Phòng Tài chính - Kế toán",
        ]

    return sorted(list(units_set))


# ---------------------------------------------------------
# 4. LỌC CHÍNH XÁC CÁN BỘ THEO ĐƠN VỊ
# ---------------------------------------------------------
def get_employees_by_unit(unit_name, all_units):
    """Chỉ lấy đúng cán bộ thuộc đơn vị chọn, chặn hoàn toàn ô đơn vị rỗng"""
    df_emp = pd.DataFrame()

    if "employees_profile" in st.session_state:
        df_ses = st.session_state["employees_profile"]
        if isinstance(df_ses, pd.DataFrame) and not df_ses.empty:
            df_emp = df_ses.copy()

    if df_emp.empty:
        supabase = init_supabase()
        if supabase:
            try:
                res = (
                    supabase.table("categories")
                    .select("content")
                    .eq("cat_key", "employees_profile")
                    .execute()
                )
                if res.data and res.data[0].get("content"):
                    df_emp = pd.DataFrame(res.data[0]["content"])
            except Exception:
                pass

    if df_emp.empty:
        return pd.DataFrame(columns=["Mã NV", "Họ và tên", "Chức vụ", "Đơn vị"])

    unit_col, name_col, id_col, pos_col = None, None, None, None

    for col in df_emp.columns:
        c_clean = str(col).strip().lower()
        if not id_col and any(
            k in c_clean
            for k in ["mã nv", "mã cb", "mã nhân viên", "ma nv", "mã cán bộ"]
        ):
            id_col = col
        elif not name_col and any(
            k in c_clean
            for k in [
                "họ và tên",
                "họ tên",
                "ho va ten",
                "tên cán bộ",
                "tên nhân viên",
            ]
        ):
            name_col = col
        elif not pos_col and any(
            k in c_clean for k in ["chức vụ", "chức danh", "vị trí"]
        ):
            pos_col = col
        elif not unit_col and any(
            k in c_clean
            for k in [
                "đơn vị",
                "khoa / phòng",
                "khoa/phòng",
                "phòng ban",
                "bộ phận",
                "khoa",
                "phòng",
            ]
        ):
            if c_clean not in ["mã đv", "mã đơn vị", "ma dv"]:
                unit_col = col

    if not name_col:
        for col in df_emp.columns:
            if "tên" in str(col).lower() or "name" in str(col).lower():
                name_col = col
                break

    if not name_col:
        return pd.DataFrame(columns=["Mã NV", "Họ và tên", "Chức vụ", "Đơn vị"])

    df_result = pd.DataFrame()
    df_result["Mã NV"] = df_emp[id_col].astype(str) if id_col else ""
    df_result["Họ và tên"] = df_emp[name_col].astype(str)
    df_result["Chức vụ"] = df_emp[pos_col].astype(str) if pos_col else ""
    df_result["Đơn vị"] = df_emp[unit_col].astype(str) if unit_col else ""

    unit_set_norm = set(normalize_str(u) for u in all_units)
    dept_prefixes = ["phòng ", "khoa ", "ban ", "tổ ", "trung tâm "]

    def is_real_person(row):
        name_norm = normalize_str(row["Họ và tên"])
        if not name_norm:
            return False
        if name_norm in unit_set_norm and name_norm != "ban giám đốc":
            return False
        for kw in dept_prefixes:
            if name_norm.startswith(kw) and name_norm != "ban giám đốc":
                return False
        return True

    df_result = df_result[df_result.apply(is_real_person, axis=1)].copy()

    if unit_name and unit_name != "Tất cả":
        target_norm = normalize_str(unit_name)

        def matches_unit(val):
            v_norm = normalize_str(val)
            if not v_norm:
                return False
            return v_norm == target_norm

        df_filtered = df_result[df_result["Đơn vị"].apply(matches_unit)].copy()
        return df_filtered

    return df_result


# ---------------------------------------------------------
# 5. TÍNH TOÁN BÁO CÁO CỘT TỔNG HỢP CHO DỮ LIỆU
# ---------------------------------------------------------
def calculate_summary_columns(df, date_cols):
    """Tự động tính các cột tổng hợp chi tiết hàng tháng cho từng nhân sự"""

    def calc_row(row):
        codes = [
            clean_and_uppercase_code(row[c]) for c in date_cols if c in row
        ]
        cong_x = sum(1 for c in codes if c in ["X", "1", "+"])
        phep_p = sum(1 for c in codes if c == "P")
        om_co = sum(1 for c in codes if c in ["ÔM", "CO"])
        cong_tac = sum(1 for c in codes if c in ["CT", "H"])
        ca_truc = sum(1 for c in codes if c == "T")
        nghi_khac = sum(1 for c in codes if c in ["TS", "KL", "R", "F"])
        tong_nghi = phep_p + om_co + nghi_khac

        return pd.Series(
            {
                "Tổng công (X)": cong_x,
                "Phép (P)": phep_p,
                "Ốm / CO": om_co,
                "Công tác (CT)": cong_tac,
                "Số ca trực (T)": ca_truc,
                "Nghỉ khác": nghi_khac,
                "Tổng ngày nghỉ": tong_nghi,
            }
        )

    summary_df = df.apply(calc_row, axis=1)
    for col in summary_df.columns:
        df[col] = summary_df[col]
    return df


# ---------------------------------------------------------
# 6. TẠO FILE EXCEL TÍCH HỢP CÔNG THỨC VÀ CỘT TỔNG HỢP CHI TIẾT
# ---------------------------------------------------------
def generate_timekeeping_excel(unit_name, month, year, df_employees):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"ChamCong_T{month}_{year}"

    # Tiêu đề
    ws.merge_cells("A1:D1")
    ws["A1"] = "BỆNH VIỆN BƯU ĐIỆN"
    ws["A1"].font = Font(name="Arial", size=10, bold=True, color="002060")

    ws.merge_cells("A2:D2")
    ws["A2"] = str(unit_name).upper()
    ws["A2"].font = Font(name="Arial", size=10, bold=True, color="002060")

    days_in_month = calendar.monthrange(year, month)[1]
    last_col_idx = 4 + days_in_month + 7
    last_col_letter = get_column_letter(last_col_idx)

    ws.merge_cells(f"A4:{last_col_letter}4")
    ws["A4"] = f"BẢNG CHẤM CÔNG THÁNG {month}/{year} CỦA {str(unit_name).upper()}"
    ws["A4"].font = Font(name="Arial", size=12, bold=True, color="002060")
    ws["A4"].alignment = Alignment(horizontal="center", vertical="center")

    headers = ["STT", "Mã NV", "Họ và tên", "Chức vụ"]
    date_cols = [f"{d:02d}/{month:02d}" for d in range(1, days_in_month + 1)]
    summary_headers = [
        "Tổng công (X)",
        "Phép (P)",
        "Ốm / CO",
        "Công tác (CT)",
        "Số ca trực (T)",
        "Nghỉ khác",
        "Tổng ngày nghỉ",
    ]

    all_headers = headers + date_cols + summary_headers

    # Định dạng Dòng 5: Tiêu đề cột
    for col_idx, h_text in enumerate(all_headers, start=1):
        cell = ws.cell(row=5, column=col_idx, value=h_text)
        cell.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )

        fill_color = "1F497D"  # Xanh navy
        if 5 <= col_idx <= 4 + days_in_month:
            d = col_idx - 4
            if is_vietnam_holiday(d, month, year):
                fill_color = "C00000"  # Đỏ Lễ
            else:
                dt = datetime(year, month, d)
                if dt.weekday() == 6:
                    fill_color = "4F81BD"  # Xanh Dương Chủ Nhật
                elif dt.weekday() == 5:
                    fill_color = "ED7D31"  # Cam Thứ 7
        elif col_idx > 4 + days_in_month:
            fill_color = "274E13"  # Xanh Lá Đậm cho Cột Tổng Hợp Báo Cáo

        cell.fill = PatternFill(
            start_color=fill_color, end_color=fill_color, fill_type="solid"
        )

    start_row = 6
    max_row = (
        start_row + len(df_employees) - 1 if not df_employees.empty else start_row + 8
    )

    first_d_letter = get_column_letter(5)
    last_d_letter = get_column_letter(4 + days_in_month)

    # Điền cán bộ & công thức tổng hợp tự động vào Excel
    for idx in range(1, (len(df_employees) if not df_employees.empty else 9) + 1):
        r = start_row + idx - 1
        if not df_employees.empty and idx <= len(df_employees):
            row_emp = df_employees.iloc[idx - 1]
            ws.cell(row=r, column=1, value=idx)
            ws.cell(row=r, column=2, value=str(row_emp.get("Mã NV") or "").strip())
            ws.cell(
                row=r, column=3, value=str(row_emp.get("Họ và tên") or "").strip()
            )
            ws.cell(
                row=r, column=4, value=str(row_emp.get("Chức vụ") or "").strip()
            )
        else:
            ws.cell(row=r, column=1, value=idx)

        ws.cell(row=r, column=1).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=2).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=3).alignment = Alignment(horizontal="left")
        ws.cell(row=r, column=4).alignment = Alignment(horizontal="left")

        # CÔNG THỨC EXCEL TỰ ĐỘNG CHO CÁC CỘT TỔNG HỢP CHI TIẾT
        c_base = 5 + days_in_month
        ws.cell(
            row=r,
            column=c_base,
            value=f'=COUNTIF({first_d_letter}{r}:{last_d_letter}{r}, "X")',
        )
        ws.cell(
            row=r,
            column=c_base + 1,
            value=f'=COUNTIF({first_d_letter}{r}:{last_d_letter}{r}, "P")',
        )
        ws.cell(
            row=r,
            column=c_base + 2,
            value=f'=COUNTIF({first_d_letter}{r}:{last_d_letter}{r}, "ÔM") + COUNTIF({first_d_letter}{r}:{last_d_letter}{r}, "CO")',
        )
        ws.cell(
            row=r,
            column=c_base + 3,
            value=f'=COUNTIF({first_d_letter}{r}:{last_d_letter}{r}, "CT") + COUNTIF({first_d_letter}{r}:{last_d_letter}{r}, "H")',
        )
        ws.cell(
            row=r,
            column=c_base + 4,
            value=f'=COUNTIF({first_d_letter}{r}:{last_d_letter}{r}, "T")',
        )
        ws.cell(
            row=r,
            column=c_base + 5,
            value=f'=COUNTIF({first_d_letter}{r}:{last_d_letter}{r}, "TS") + COUNTIF({first_d_letter}{r}:{last_d_letter}{r}, "KL") + COUNTIF({first_d_letter}{r}:{last_d_letter}{r}, "R")',
        )

        p_col_let = get_column_letter(c_base + 1)
        om_col_let = get_column_letter(c_base + 2)
        nk_col_let = get_column_letter(c_base + 5)
        ws.cell(
            row=r,
            column=c_base + 6,
            value=f"={p_col_let}{r}+{om_col_let}{r}+{nk_col_let}{r}",
        )

    # Tô màu Nền Ngày Lễ, Thứ 7, Chủ Nhật (Chủ Nhật = XANH DƯƠNG NHẠT #D9E1F2)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    holiday_fill = PatternFill(
        start_color="FADBD8", end_color="FADBD8", fill_type="solid"
    )
    sun_fill = PatternFill(
        start_color="D9E1F2", end_color="D9E1F2", fill_type="solid"
    )
    sat_fill = PatternFill(
        start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"
    )
    summary_fill = PatternFill(
        start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"
    )

    for r in range(6, max_row + 1):
        for c in range(1, last_col_idx + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = thin_border
            cell.font = Font(name="Arial", size=9)

            if 5 <= c <= 4 + days_in_month:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                d = c - 4
                if is_vietnam_holiday(d, month, year):
                    cell.fill = holiday_fill
                else:
                    dt = datetime(year, month, d)
                    if dt.weekday() == 6:
                        cell.fill = sun_fill
                    elif dt.weekday() == 5:
                        cell.fill = sat_fill
            elif c > 4 + days_in_month:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.fill = summary_fill
                cell.font = Font(name="Arial", size=9, bold=True)

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 25
    ws.column_dimensions["D"].width = 20
    for c in range(5, 5 + days_in_month):
        ws.column_dimensions[get_column_letter(c)].width = 7
    for c in range(5 + days_in_month, last_col_idx + 1):
        ws.column_dimensions[get_column_letter(c)].width = 13

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


# ---------------------------------------------------------
# 7. GIAO DIỆN PHÂN HỆ CHẤM CÔNG VÀ BÁO CÁO TỔNG HỢP
# ---------------------------------------------------------
def render_timekeeping_management():
    st.markdown(
        "<h3 style='color: #003366;'>📅 Chấm công - Ca trực - Phân lịch</h3>",
        unsafe_allow_html=True,
    )

    supabase = init_supabase()
    all_units = load_system_units(supabase)

    col_f1, col_f2, col_f3 = st.columns([2, 2, 4])
    with col_f1:
        selected_month = st.selectbox(
            "Chọn Tháng:", list(range(1, 13)), index=8
        )
    with col_f2:
        selected_year = st.selectbox("Chọn Năm:", [2025, 2026, 2027], index=1)
    with col_f3:
        selected_unit = st.selectbox("Chọn Đơn vị / Khoa / Phòng:", all_units)

    df_unit_emp = get_employees_by_unit(selected_unit, all_units)

    st.markdown("---")
    col_act1, col_act2 = st.columns([3, 4])

    with col_act1:
        excel_data = generate_timekeeping_excel(
            selected_unit, selected_month, selected_year, df_unit_emp
        )
        file_name_clean = selected_unit.replace(" ", "_")
        st.download_button(
            label=f"📥 Tải File Excel Mẫu Chấm Công ({len(df_unit_emp)} Cán bộ)",
            data=excel_data,
            file_name=(
                f"Mau_Cham_Cong_{file_name_clean}_T{selected_month}_{selected_year}.xlsx"
            ),
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary",
        )

    with col_act2:
        uploaded_file = st.file_uploader(
            "📤 Upload Tệp Bảng Chấm Công (.xlsx)",
            type=["xlsx", "xls"],
            key="up_tk_file",
        )

    if uploaded_file:
        try:
            df_up = pd.read_excel(uploaded_file, skiprows=4).fillna("")
            for col in df_up.columns:
                if "/" in str(col):
                    df_up[col] = df_up[col].apply(clean_and_uppercase_code)
            st.session_state[
                f"tk_{selected_unit}_{selected_month}_{selected_year}"
            ] = df_up
            st.toast("✅ Đã tải lên thành công bảng chấm công!", icon="🎉")
            st.success("✅ Đã cập nhật thành công bảng chấm công từ file Excel!")
        except Exception as e:
            st.error(f"Lỗi đọc file chấm công: {e}")

    st.markdown("---")
    st.write(
        f"📋 **Bảng Chấm công Tháng {selected_month}/{selected_year} -"
        f" {selected_unit}:**"
    )

    days_in_month = calendar.monthrange(selected_year, selected_month)[1]
    tk_key = f"tk_{selected_unit}_{selected_month}_{selected_year}"

    if tk_key in st.session_state:
        df_tk_display = st.session_state[tk_key]
    else:
        date_cols_list = [
            f"{d:02d}/{selected_month:02d}" for d in range(1, days_in_month + 1)
        ]
        cols = (
            ["STT", "Mã NV", "Họ và tên", "Chức vụ"]
            + date_cols_list
            + [
                "Tổng công (X)",
                "Phép (P)",
                "Ốm / CO",
                "Công tác (CT)",
                "Số ca trực (T)",
                "Nghỉ khác",
                "Tổng ngày nghỉ",
            ]
        )

        if not df_unit_emp.empty and "Họ và tên" in df_unit_emp.columns:
            data_rows = []
            for idx, (_, r) in enumerate(df_unit_emp.iterrows(), start=1):
                row_dict = {
                    "STT": idx,
                    "Mã NV": str(r.get("Mã NV") or "").strip(),
                    "Họ và tên": str(r.get("Họ và tên") or "").strip(),
                    "Chức vụ": str(r.get("Chức vụ") or "").strip(),
                }
                for d in range(1, days_in_month + 1):
                    row_dict[f"{d:02d}/{selected_month:02d}"] = ""
                data_rows.append(row_dict)
            df_tk_display = pd.DataFrame(data_rows)
        else:
            df_tk_display = pd.DataFrame(columns=cols)
            st.info(
                f"ℹ️ Đơn vị **{selected_unit}** hiện chưa có hồ sơ cán bộ đăng"
                " ký trong hệ thống. Bạn có thể thêm dòng trực tiếp trên bảng"
                " bên dưới hoặc nhập danh sách tại mục **4. Quản lý Hồ sơ Cán"
                " bộ**."
            )

    date_cols = [c for c in df_tk_display.columns if "/" in str(c)]

    # 1. TỰ ĐỘNG CHUYỂN TOÀN BỘ KÝ HIỆU SANG CHỮ IN HOA
    for col in date_cols:
        df_tk_display[col] = df_tk_display[col].apply(clean_and_uppercase_code)

    # 2. TÍNH LẠI CÁC CỘT TỔNG HỢP CHI TIẾT BÊN PHẢI
    df_tk_display = calculate_summary_columns(df_tk_display, date_cols)

    # 3. PHÁT HIỆN VÀ CẢNH BÁO KÝ TỰ LẠ
    invalid_entries = []
    valid_codes_set = set(VALID_CODES_MAP.values())

    for idx, row in df_tk_display.iterrows():
        emp_name = str(row.get("Họ và tên") or f"Dòng {idx+1}").strip()
        for col in date_cols:
            val = str(row[col]).strip()
            if val and val not in valid_codes_set:
                invalid_entries.append(
                    f"- **{emp_name}** (Ngày **{col}**): Ký tự lạ `{val}`"
                )

    if invalid_entries:
        st.warning(
            "⚠️ **CẢNH BÁO: PHÁT HIỆN KÝ TỰ CHẤM CÔNG LẠ KHÔNG ĐÚNG QUY"
            " ĐỊNH!**\n"
            + "\n".join(invalid_entries[:8])
            + ("\n... và các ô khác." if len(invalid_entries) > 8 else "")
            + "\n\n👉 *Ký hiệu chuẩn:* **X** (Công), **P** (Phép), **ÔM**"
            " (Ốm), **CO** (Con ốm), **CT** (Công tác), **T** (Trực), **TS**"
            " (Thai sản), **KL** (Không lương), **R** (Việc riêng), **H** (Học)."
        )

    # Định dạng màu sắc cột (Chủ Nhật = Xanh Dương Nhạt)
    def highlight_days(col):
        col_name = str(col.name)
        if "/" in col_name:
            try:
                day_num = int(col_name.split("/")[0])
                if is_vietnam_holiday(day_num, selected_month, selected_year):
                    return [
                        "background-color: #FADBD8; color: #78281F;"
                        " font-weight: bold"
                    ] * len(col)
                dt = datetime(selected_year, selected_month, day_num)
                if dt.weekday() == 6:  # Chủ Nhật
                    return [
                        "background-color: #D9E1F2; color: #1F497D;"
                        " font-weight: bold"
                    ] * len(col)
                elif dt.weekday() == 5:  # Thứ 7
                    return [
                        "background-color: #FFF2CC; color: #9C6500"
                    ] * len(col)
            except Exception:
                pass
        elif any(
            k in col_name
            for k in [
                "Tổng công",
                "Phép",
                "Ốm",
                "Công tác",
                "trực",
                "khác",
                "nghỉ",
            ]
        ):
            return [
                "background-color: #E2EFDA; color: #274E13; font-weight: bold"
            ] * len(col)
        return [""] * len(col)

    styled_df = df_tk_display.style.apply(highlight_days, axis=0)

    # Disable không cho sửa trực tiếp các cột Tổng hợp
    disabled_cols = [
        "STT",
        "Mã NV",
        "Họ và tên",
        "Chức vụ",
        "Tổng công (X)",
        "Phép (P)",
        "Ốm / CO",
        "Công tác (CT)",
        "Số ca trực (T)",
        "Nghỉ khác",
        "Tổng ngày nghỉ",
    ]

    edited_tk = st.data_editor(
        styled_df,
        use_container_width=True,
        num_rows="dynamic",
        disabled=[c for c in disabled_cols if c in df_tk_display.columns],
        key=f"editor_{tk_key}",
    )

    if st.button("💾 Lưu Bảng Chấm Công", type="primary"):
        # CHUẨN HÓA LẠI TRƯỚC KHI LƯU DỮ LIỆU
        for col in date_cols:
            if col in edited_tk.columns:
                edited_tk[col] = edited_tk[col].apply(clean_and_uppercase_code)

        edited_tk = calculate_summary_columns(edited_tk, date_cols)
        st.session_state[tk_key] = edited_tk
        st.toast(
            "✅ Đã lưu và cập nhật tổng hợp dữ liệu chấm công thành công!",
            icon="💾",
        )
        st.success("✅ Đã lưu thành công dữ liệu chấm công!")
